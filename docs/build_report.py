"""Build docs/MLOps_Assignment1_Report.docx from project artifacts.

Screenshots are picked up from screenshots/ by filename; missing ones become placeholders.
Run: python docs/build_report.py
"""
import json
import re
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FIGS = ROOT / "reports" / "figures"
SHOTS = ROOT / "screenshots"
OUT = ROOT / "docs" / "MLOps_Assignment1_Report.docx"
ARCH = ROOT / "docs" / "architecture.png"
REPO = "https://github.com/avdurgesh/MlOps-heart-diseae-durgesh-verma-AIMLCZG523"
METRICS = json.loads((ROOT / "models" / "metrics.json").read_text())


def draw_architecture() -> None:
    fig, ax = plt.subplots(figsize=(11, 4.6))
    ax.axis("off")
    ax.set_xlim(0, 11)
    ax.set_ylim(0, 4.6)

    def box(x, y, w, h, text, color):
        ax.add_patch(plt.Rectangle((x, y), w, h, fc=color, ec="#333333", lw=1.2, zorder=2))
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, zorder=3)

    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1),
                    arrowprops=dict(arrowstyle="->", lw=1.4, color="#333333"), zorder=1)

    top, bot = 3.0, 0.8
    box(0.1, top, 1.8, 1.0, "UCI Heart\nDisease data", "#dbe9f6")
    box(2.3, top, 1.9, 1.0, "src/data.py\nclean + EDA", "#dbe9f6")
    box(4.6, top, 2.1, 1.0, "src/train.py\nGridSearchCV + 5-fold CV", "#fde9c9")
    box(7.1, top, 1.7, 1.0, "MLflow\nexperiments", "#fde9c9")
    box(9.1, top, 1.8, 1.0, "models/\nmodel.joblib", "#e2f0d9")
    box(0.1, bot, 1.8, 1.0, "GitHub Actions\nlint, test, train", "#eeeeee")
    box(2.3, bot, 1.9, 1.0, "FastAPI app\n/predict /health /metrics", "#e2f0d9")
    box(4.6, bot, 1.7, 1.0, "Docker\nimage", "#e2f0d9")
    box(6.7, bot, 2.1, 1.0, "Kubernetes (Minikube)\nDeployment + Service\n+ Ingress", "#e2f0d9")
    box(9.1, bot, 1.8, 1.0, "Prometheus\n+ Grafana", "#f4d6d6")
    for a, b in [(1.9, 2.3), (4.2, 4.6), (6.7, 7.1)]:
        arrow(a, top + 0.5, b, top + 0.5)
    arrow(8.8, top + 0.5, 9.1, top + 0.5)
    arrow(10.0, top, 3.25, bot + 1.0)
    arrow(4.2, bot + 0.5, 4.6, bot + 0.5)
    arrow(6.3, bot + 0.5, 6.7, bot + 0.5)
    arrow(8.8, bot + 0.5, 9.1, bot + 0.5)
    arrow(1.0, bot + 1.0, 4.6, top, )
    ax.text(5.5, 2.35, "model loaded at startup", fontsize=8, style="italic", ha="center")
    fig.tight_layout()
    fig.savefig(ARCH, dpi=170)
    plt.close(fig)


def pct(x: float) -> str:
    return f"{x:.3f}"


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(1)
    s.top_margin = s.bottom_margin = Inches(0.9)


def h(text, level=1):
    return doc.add_heading(text, level=level)


def p(text, bold=False, italic=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold, run.italic = bold, italic
    return para


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    para.paragraph_format.left_indent = Inches(0.25)


def caption(text):
    c = doc.add_paragraph()
    r = c.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER


FIG = [0]


def next_caption(text: str) -> str:
    FIG[0] += 1
    return f"Figure {FIG[0]} - {re.sub(r'^Figure \d+ - ', '', text)}"


def picture(path: Path, text: str, width=6.0):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption(next_caption(text))


def shot(name: str, text: str, width=6.0):
    path = SHOTS / name
    if path.exists():
        picture(path, text, width)
    else:
        para = doc.add_paragraph()
        r = para.add_run(f"[Insert screenshot: screenshots/{name} - {text}]")
        r.bold = True
        r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(rows, header, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, c in enumerate(header):
        t.rows[0].cells[i].text = c
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            cells[i].text = str(c)
    doc.add_paragraph()


draw_architecture()

# ---------------- Title page ----------------
for _ in range(6):
    doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Machine Learning Operations (MLOps) - AIMLCZG523")
r.font.size = Pt(16)
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Assignment 01\nEnd-to-End ML Model Development, CI/CD and Production Deployment")
r.bold = True
r.font.size = Pt(22)
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.add_run("Heart Disease Risk Prediction (UCI Heart Disease Dataset)").font.size = Pt(14)
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.add_run("Student: Durgesh Verma\n").bold = True
t.add_run(f"Code repository: {REPO}")
doc.add_page_break()

# ---------------- 1 Introduction ----------------
h("1. Introduction and objectives")
p("This report describes an end-to-end MLOps solution that predicts the risk of heart disease from patient "
  "health data. The classifier is trained and tracked reproducibly, tested and trained automatically in a "
  "CI pipeline, packaged as a Docker image, deployed to Kubernetes, and monitored with Prometheus and Grafana.")
p("Objectives:", bold=True)
bullets([
    "Acquire, clean and explore the UCI Heart Disease (Cleveland) dataset.",
    "Build and compare at least two classifiers with cross-validation and proper metrics.",
    "Track every experiment (parameters, metrics, plots, models) with MLflow.",
    "Package the model together with its preprocessing so serving is fully reproducible.",
    "Automate linting, unit tests and training with GitHub Actions.",
    "Serve the model through a containerised REST API and deploy it on Kubernetes.",
    "Log API requests and expose metrics for monitoring.",
])
h("Repository layout", 2)
table([
    ["src/", "data loading and cleaning, EDA, features, training, inference"],
    ["app/", "FastAPI service (/predict, /health, /metrics)"],
    ["tests/", "pytest unit tests for data, model and API (14 tests)"],
    ["models/", "saved pipeline (model.joblib) and metrics.json"],
    ["k8s/", "Kubernetes Deployment, Service and Ingress manifests"],
    ["monitoring/", "Prometheus config and provisioned Grafana dashboard"],
    [".github/workflows/", "GitHub Actions CI workflow"],
], ["Path", "Content"])

# ---------------- 2 Setup ----------------
h("2. Setup and installation")
p("The project runs from a clean checkout using only requirements.txt (Python 3.12).")
code("git clone " + REPO + ".git\ncd MlOps-heart-diseae-durgesh-verma-AIMLCZG523\n"
     "python3 -m venv .venv && source .venv/bin/activate   # Windows: .\\.venv\\Scripts\\Activate.ps1\n"
     "pip install -r requirements.txt")
p("Pipeline commands:")
code("python -m src.data      # download + clean dataset\n"
     "python -m src.eda       # EDA figures\n"
     "python -m src.train     # tune, evaluate, log to MLflow, save model\n"
     "pytest                  # unit tests\n"
     "flake8 src tests app    # lint")
p("The cleaned dataset is committed to the repository, so training also works when the machine has no "
  "internet access; src.data re-downloads the raw file from the UCI repository when needed.")
shot("vm_data_eda.png", "Output of src.data and src.eda", 6.2)
shot("vm_train_tests.png", "Output of src.train, pytest and src.predict", 6.2)

# ---------------- 3 Data and EDA ----------------
doc.add_page_break()
h("3. Data acquisition and exploratory data analysis")
p("Source: UCI Machine Learning Repository, Heart Disease dataset (Cleveland subset): 303 patients, 13 input "
  "features and one target. Missing values are encoded as '?' in the raw file (4 values in 'ca' and 2 in 'thal').")
h("3.1 Cleaning and preprocessing", 2)
bullets([
    "Duplicates removed; the file contains none, but the step keeps the pipeline safe for new data.",
    "The original target has values 0-4 (disease severity); it is collapsed to a binary label "
    "(0 = no disease, 1 = disease), matching the problem statement.",
    "Missing 'ca' and 'thal' values are imputed with the column mode (categorical features).",
    "Categorical codes are cast to integers; no missing values remain after cleaning (303 rows x 14 columns).",
])
h("3.2 Findings", 2)
bullets([
    "Class balance: 164 patients without disease and 139 with disease (54% / 46%), so the classes are "
    "roughly balanced and no resampling is needed.",
    "Strongest positive correlations with the target: thal (0.52), ca (0.46), exang (0.43), oldpeak (0.42) "
    "and cp (0.41). Strongest negative: thalach (-0.42), i.e. a lower maximum heart rate goes with disease.",
    "Weak signals: fbs (0.03) and chol (0.09), and cholesterol also contains a few high outliers.",
    "Numeric features are on very different scales (chol up to 564, oldpeak up to 6.2), which motivates "
    "standardisation for the linear model.",
])
picture(FIGS / "class_balance.png", "Figure 1 - Class balance of the binary target", 3.4)
picture(FIGS / "numeric_histograms.png", "Figure 2 - Distribution of numeric features by class", 6.2)
picture(FIGS / "correlation_heatmap.png", "Figure 3 - Correlation heatmap", 5.4)
picture(FIGS / "categorical_vs_target.png", "Figure 4 - Categorical features versus the target", 6.2)

# ---------------- 4 Modelling ----------------
doc.add_page_break()
h("4. Feature engineering and model development")
h("4.1 Feature pipeline", 2)
p("All preprocessing lives inside a single scikit-learn Pipeline (src/features.py), so exactly the same "
  "transformations are applied at training time and at serving time:")
bullets([
    "Numeric features (age, trestbps, chol, thalach, oldpeak): StandardScaler.",
    "Categorical features (sex, cp, fbs, restecg, exang, slope, ca, thal): OneHotEncoder with "
    "handle_unknown='ignore', so an unseen category cannot crash the API (covered by a unit test).",
])
h("4.2 Models, tuning and evaluation", 2)
p("Two classifiers were trained on a stratified 80/20 split (random_state=42): Logistic Regression as an "
  "interpretable baseline and Random Forest as a non-linear model. Hyper-parameters were tuned with "
  "GridSearchCV using 5-fold stratified cross-validation and ROC-AUC as the selection metric. The tuned "
  "pipeline was then evaluated with cross-validation (accuracy, precision, recall, ROC-AUC) on the training "
  "data and once on the held-out test set.")
table([
    ["Logistic Regression", "C in {0.01, 0.1, 1, 10}", "C = 1"],
    ["Random Forest", "n_estimators {100, 300}, max_depth {3, 5, None}, min_samples_leaf {1, 3}",
     "n_estimators = 100, max_depth = 5, min_samples_leaf = 3"],
], ["Model", "Search grid", "Best parameters"])
rows = []
for name, label in [("logistic_regression", "Logistic Regression"), ("random_forest", "Random Forest")]:
    m = METRICS["results"][name]
    rows.append([label, pct(m["cv_accuracy"]), pct(m["cv_precision"]), pct(m["cv_recall"]),
                 pct(m["cv_roc_auc"]), pct(m["test_accuracy"]), pct(m["test_recall"]), pct(m["test_roc_auc"])])
table(rows, ["Model", "CV acc", "CV prec", "CV recall", "CV ROC-AUC", "Test acc", "Test recall", "Test ROC-AUC"])
p("Model selection: both models reach almost the same cross-validated ROC-AUC (0.902 versus 0.901). Logistic "
  "Regression has higher cross-validated accuracy, precision and recall, a higher test ROC-AUC (0.966 versus "
  "0.947), is simpler and is easier to explain, so it was selected as the final model. For a screening "
  "use case recall matters, and it reaches 0.93 on the test set. With only 303 samples the test set holds "
  "61 patients, so these test numbers have wide uncertainty and the cross-validated figures are the more "
  "reliable comparison.")

# ---------------- 5 MLflow ----------------
h("5. Experiment tracking with MLflow")
p("Each model is trained inside an MLflow run (experiment 'heart-disease-classification'). A run logs:")
bullets([
    "Parameters: model type, best hyper-parameters and number of features.",
    "Metrics: cross-validated accuracy, precision, recall and ROC-AUC, plus the same four on the test set.",
    "Artifacts: ROC curve, confusion matrix and the fitted pipeline as an MLflow model with an input example "
    "(and therefore an inferred signature).",
])
p("The UI is started with: mlflow ui --backend-store-uri ./mlruns.")
shot("mlflow_experiments.png", "Figure 5 - MLflow experiment with both runs")
shot("mlflow_compare.png", "Figure 6 - Run comparison")
shot("mlflow_run_artifacts.png", "Figure 7 - Artifacts of the selected run (ROC curve, confusion matrix, model)")

# ---------------- 6 Packaging ----------------
doc.add_page_break()
h("6. Model packaging and reproducibility")
bullets([
    "The final pipeline (preprocessing + Logistic Regression) is saved with joblib as models/model.joblib "
    "and is also stored as an MLflow model. metrics.json records the winning model and its scores.",
    "requirements.txt pins version ranges for the full development environment; requirements-api.txt pins "
    "exact versions (scikit-learn 1.9.1, numpy, pandas, scipy) for the serving image, because a pickled "
    "model must be loaded with the same scikit-learn version it was trained with.",
    "A fixed random_state (42) and a deterministic split make training repeatable; the CI pipeline proves "
    "that training runs from a clean environment.",
    "src/predict.py is the single inference entry point used by both the command line and the API.",
])

# ---------------- 7 CI/CD ----------------
h("7. CI/CD pipeline and automated testing")
p("Unit tests (pytest, 14 tests) cover three areas:")
bullets([
    "Data: cleaning removes duplicates, binarises the target, fills missing values and does not mutate its input.",
    "Model: shape of the processed dataset, preprocessor output without NaNs, a fitted pipeline beats a "
    "sanity threshold, valid prediction and confidence range, and robustness to an unseen category.",
    "API: /health, /predict with valid input, rejection of invalid or missing fields (HTTP 422) and the "
    "Prometheus /metrics endpoint.",
])
p("The GitHub Actions workflow (.github/workflows/ci.yml) runs on every push and pull request on ubuntu-latest "
  "with Python 3.12: install requirements, lint with flake8, run pytest with coverage, prepare the data, "
  "train the models, and upload models, MLflow runs, the training log, test results and coverage as the "
  "'ml-artifacts' artifact. The training step runs under bash with pipefail, so the job fails if "
  "training fails even though the log is piped through tee. Any lint, test or training error fails the pipeline.")
shot("ci_run.png", "Successful GitHub Actions runs (CI #1 and CI #2)")

# ---------------- 8 Docker ----------------
doc.add_page_break()
h("8. Model containerisation")
p("The model is served by a FastAPI application (app/main.py):")
table([
    ["POST /predict", "JSON with the 13 features; returns prediction (0/1), confidence and probability_disease"],
    ["GET /health", "liveness/readiness check; reports whether the model is loaded"],
    ["GET /metrics", "Prometheus metrics"],
    ["GET /docs", "interactive Swagger UI"],
], ["Endpoint", "Description"])
p("Inputs are validated with Pydantic (type and range constraints), so malformed requests return HTTP 422 "
  "instead of reaching the model. The Dockerfile uses python:3.12-slim, installs only requirements-api.txt, "
  "copies the source and the saved model, runs as a non-root user and defines a HEALTHCHECK. The resulting "
  "image is about 470 MB.")
code("docker build -t heart-disease-api:1.0 .\n"
     "docker run -d -p 8000:8000 heart-disease-api:1.0\n"
     'curl -X POST http://localhost:8000/predict -H "Content-Type: application/json" \\\n'
     "  -d '{\"age\":63,\"sex\":1,\"cp\":1,\"trestbps\":145,\"chol\":233,\"fbs\":1,\"restecg\":2,"
     "\"thalach\":150,\"exang\":0,\"oldpeak\":2.3,\"slope\":3,\"ca\":0,\"thal\":6}'")
p('Response: {"prediction": 0, "confidence": 0.834, "probability_disease": 0.166}')
shot("docker_build_run.png", "Figure 9 - Docker build, run and a /predict request")
shot("swagger.png", "Figure 10 - Swagger UI for the API")

# ---------------- 9 Deployment ----------------
h("9. Production deployment on Kubernetes")
p("The image is deployed to a Minikube cluster (Docker driver) using the manifests in k8s/:")
bullets([
    "Namespace heart-disease and a Deployment with 2 replicas, readiness and liveness probes on /health, "
    "and CPU/memory requests and limits.",
    "A Service of type LoadBalancer on port 80 forwarding to container port 8000.",
    "An nginx Ingress routing / to the Service.",
])
code("minikube start -p heart-disease --driver=docker\n"
     "minikube -p heart-disease image load heart-disease-api:1.0\n"
     "kubectl apply -f k8s/deployment.yaml\n"
     "minikube -p heart-disease addons enable ingress\n"
     "kubectl apply -f k8s/ingress.yaml\n"
     "curl http://$(minikube -p heart-disease ip)/health")
p("Verification: both pods reach Running and the Ingress receives the address of the Minikube node. "
  "/health and /predict called through the Ingress return the same results as the plain Docker container "
  "(prediction 0, confidence 0.834). The LoadBalancer Service keeps an EXTERNAL-IP of <pending> because "
  "Minikube has no cloud load balancer; the Ingress is the access path.")
shot("k8s_resources.png", "kubectl get pods, svc and ingress, the /metrics counters and the request logs of the pods")

# ---------------- 10 Monitoring ----------------
doc.add_page_break()
h("10. Monitoring and logging")
p("Logging: a middleware logs every request (method, path, status, latency in ms) and every prediction "
  "(label and confidence) to standard output, which is collected by 'docker logs' / 'kubectl logs'.")
p("Metrics: the API exposes Prometheus metrics at /metrics:")
bullets([
    "api_requests_total{method, path, status} - request counter.",
    "api_request_latency_seconds{path} - latency histogram.",
    "model_predictions_total{label} - predictions served per class.",
])
p("docker compose starts the API, Prometheus (scraping the API every 15 s) and Grafana. The Prometheus data "
  "source and a 'Heart Disease API' dashboard are provisioned from files in monitoring/grafana/, so no manual "
  "setup is needed. The dashboard shows total requests and predictions, predictions by class, request rate "
  "per path and p95 latency.")
shot("prometheus_targets.png", "Figure 13 - Prometheus target heart-disease-api is UP")
shot("grafana_dashboard.png", "Figure 14 - Grafana dashboard")

# ---------------- 11 Architecture ----------------
h("11. Architecture")
picture(ARCH, "Figure 15 - End-to-end architecture", 6.5)
p("Data flows from the UCI source through cleaning and EDA into training. Training logs to MLflow and "
  "produces the packaged pipeline, which the FastAPI service loads at start-up. The service is built into a "
  "Docker image, deployed to Kubernetes, and observed by Prometheus and Grafana. GitHub Actions validates "
  "every change by linting, testing and re-training.")

# ---------------- 12 Challenges ----------------
h("12. Issues encountered and how they were solved")
bullets([
    "Pickle compatibility: the serving image pins the exact scikit-learn version used for training.",
    "Deprecated scikit-learn argument: the 'penalty' parameter of LogisticRegression was removed from the "
    "grid and the default L2 regularisation is used.",
    "minikube image load failed on Windows (missing wmic); the image is built directly inside Minikube's "
    "Docker daemon instead.",
    "Port 3000 was already in use on the development machine, so Grafana is exposed on port 3001.",
    "Docker Desktop BuildKit reported 'changes out of order' while rebuilding through compose; the stack was "
    "started from the already built image.",
    "Pipeline robustness: tee would hide a failing training step, so the step runs with 'shell: bash' (pipefail).",
])

# ---------------- 13 Conclusion ----------------
h("13. Conclusion")
p("The project delivers a reproducible pipeline: a clean dataset and EDA, two tuned and cross-validated "
  "models tracked in MLflow, a packaged preprocessing + model pipeline, an automated CI pipeline that lints, "
  "tests and trains, a containerised API deployed on Kubernetes behind an Ingress, and request logging with "
  "Prometheus and Grafana monitoring. The selected Logistic Regression model reaches a cross-validated ROC-AUC "
  "of 0.90 and a test ROC-AUC of 0.97.")
p("Possible improvements: a larger multi-site dataset, probability calibration, a model registry with "
  "automated promotion, image push to a container registry in CI (continuous delivery), and drift monitoring "
  "of the input features.")
p("Disclaimer: this model is a learning exercise and is not a medical device.", italic=True)
h("Access instructions", 2)
bullets([
    f"Code repository: {REPO}",
    "API (Docker): http://localhost:8000 (docs at /docs).",
    "Kubernetes: through the Ingress at the Minikube IP ('minikube ip'), or "
    "'kubectl -n heart-disease port-forward svc/heart-disease-api 8080:80'.",
    "Grafana: http://localhost:3001 (admin/admin), Prometheus: http://localhost:9090, MLflow: http://localhost:5001.",
])

OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f"Saved {OUT}")
